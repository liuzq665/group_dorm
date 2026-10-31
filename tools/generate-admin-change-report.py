from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "基本资料与角色管理改进总结及工作留痕.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17324D"
TEAL = "176B70"
PALE = "EAF2F4"
LIGHT = "F3F6F8"
MID = "D8E1E7"
TEXT = "24313B"
MUTED = "5D6B75"


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    later = [props.index(node) for name in ("tcMar", "textDirection", "tcFitText", "vAlign", "hideMark") if (node := props.find(qn(f"w:{name}"))) is not None]
    props.insert(min(later) if later else len(props), shading)


def set_cell_margins(cell, top=95, start=110, bottom=95, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        valign = tc_pr.find(qn("w:vAlign"))
        tc_pr.insert(tc_pr.index(valign) if valign is not None else len(tc_pr), margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell_text(cell, text, *, bold=False, color=TEXT, size=9.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths=None, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=font_size)
        set_cell_shading(cell, NAVY)
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            set_cell_text(cells[idx], value, size=font_size)
            if row_index % 2 == 1:
                set_cell_shading(cells[idx], LIGHT)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_p(doc, text, *, style=None, bold_prefix=None, color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        first.bold = True
        rest = p.add_run(text[len(bold_prefix):])
        runs = (first, rest)
    else:
        runs = (p.add_run(text),)
    for run in runs:
        run.font.name = "Microsoft YaHei"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    paragraph.add_run(" 页")


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
section.header_distance = Inches(0.32)
section.footer_distance = Inches(0.32)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.18
for style_name, size, color in (("Title", 27, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 13, TEAL), ("Heading 3", 11, NAVY)):
    st = doc.styles[style_name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
    st.paragraph_format.space_after = Pt(5)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header_run = header.add_run("宿舍事务平台  |  管理端改进留痕")
header_run.font.name = "Microsoft YaHei"
header_run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
header_run.font.size = Pt(8)
header_run.font.color.rgb = RGBColor.from_string(MUTED)
add_page_number(section.footer.paragraphs[0])

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(95)
p.paragraph_format.space_after = Pt(12)
r = p.add_run("宿舍事务平台")
r.bold = True
r.font.name = "Microsoft YaHei"
r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(16)
r.font.color.rgb = RGBColor.from_string(TEAL)

p = doc.add_paragraph(style="Title")
p.paragraph_format.space_after = Pt(12)
p.add_run("基本资料与角色管理\n改进总结及工作留痕")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
r = p.add_run("从最初的后台列表诉求，到当前数据模型、业务流程与页面交互")
r.font.size = Pt(12)
r.font.color.rgb = RGBColor.from_string(MUTED)

doc.add_paragraph()
cover = doc.add_table(rows=4, cols=2)
cover.alignment = WD_TABLE_ALIGNMENT.LEFT
cover.autofit = False
cover.style = "Table Grid"
cover_data = [
    ("文档目的", "按需求演进顺序记录基本资料与角色管理的修改，并给出当前实现证据"),
    ("项目范围", "系统管理员端：基本资料（楼栋、宿舍、床位）与角色管理（学生、宿管、维修人员）"),
    ("证据来源", "当前 SQL 数据库快照、前后端代码、浏览器测试脚本及本对话中的明确需求"),
    ("整理日期", "2026 年 10 月 30 日"),
]
for i, (label, value) in enumerate(cover_data):
    set_cell_text(cover.rows[i].cells[0], label, bold=True, color=NAVY, size=9.5)
    set_cell_shading(cover.rows[i].cells[0], PALE)
    set_cell_text(cover.rows[i].cells[1], value, size=9.5)
    cover.rows[i].cells[0].width = Inches(1.25)
    cover.rows[i].cells[1].width = Inches(5.5)

doc.add_page_break()

add_heading(doc, "阅读说明", 1)
add_p(doc, "本报告比较的是对话中最早提出的系统管理员基本资料与账号/角色管理诉求，和当前工作区可核对的实现。对话里曾多次补充、纠正同一交互，本报告将后续明确要求视为最终规则。")
add_p(doc, "“已落地”表示当前代码或数据库结构中可以直接核对到相应实现；“有测试留痕”表示仓库里存在相应浏览器测试脚本或历史结果；它不自动表示该脚本在本次整理时重新执行通过。")

add_heading(doc, "目录", 1)
for item in [
    "1. 核心变化一览",
    "2. 按需求演进的工作留痕",
    "3. 基本资料：楼栋、宿舍、床位",
    "4. 角色管理：学生、宿管、维修人员",
    "5. 数据模型与业务一致性",
    "6. 页面布局与操作体验的反复修正",
    "7. 当前验证证据与边界",
    "8. 关键代码与数据文件索引",
]:
    add_bullet(doc, item)

add_heading(doc, "1. 核心变化一览", 1)
add_p(doc, "最初的关注点集中在系统管理员能否查看、维护账号与基础信息；随着实际使用逐步明确，功能扩展为宿舍资源管理和按角色维护个人资料的业务后台。当前实现把住宿位置、人员身份、学籍信息、员工范围与账号状态分别纳入相应的数据关系，并为列表查询、详情、编辑和分页补齐了完整交互。")
add_table(doc, ["模块", "最初关注", "当前形态"], [
    ("基本资料", "查看楼栋、宿舍、床位，能进行基础维护", "楼栋汇总和楼下宿舍列表；宿舍床位与学生信息；床位位置、占用和管理状态；筛选级联、分页、床位入住调整"),
    ("角色管理", "学生/宿管/维修人员列表与账号管理", "一级“角色管理”含三个角色子页；角色档案字段、角色专属详情、添加编辑、启停、删除/批量删除、范围配置"),
    ("数据关系", "增加学生学号和个人资料", "独立学生资料表；员工角色资料表；person 身份主键；学院/专业/班级字典关系；住宿分配历史表及查询视图"),
    ("页面体验", "查询和操作页面可以使用", "范围筛选、级联下拉、记录分页与页码跳转；对齐操作列、控制换行、适配窄屏并减少多余滚动"),
], widths=[1.0, 2.2, 3.7])

add_heading(doc, "2. 按需求演进的工作留痕", 1)
timeline = [
    ("阶段 1：从共享账号视角转为角色资料视角", "最初围绕账号列表、角色列表和学生学号提出修改；随后明确学生学号是学生资料，不应塞进通用 user 字段。", "新增 student 专属资料；后续进一步提出宿管、维修人员、系统管理员各自拥有个人资料。", "当前数据库有 person 与 student、dorm_manager_profile、repairer_profile、system_admin_profile 的一对一扩展关系；通用账号表已不作为当前角色资料主表。"),
    ("阶段 2：明确学生和员工的基础资料边界", "学生信息原先不足以支撑住宿管理；员工档案也缺少工号、个人资料和业务范围。", "学生补充学号、姓名、性别、生日、电话、账号状态、班级与学籍状态；员工档案补充工号、姓名、性别、生日、电话、入职时间、启停和在岗状态。", "当前 SQL 中 student.student_no 为 char(10) 且有格式约束；DM/RP/SA 工号有统一格式约束；角色资料含创建/更新时间。"),
    ("阶段 3：角色菜单与角色详情", "管理功能容易把账号和角色混在一起，角色档案详情内容未定义。", "改为“角色管理”一级菜单，学生、宿管、维修人员为二级页面；详情按角色展示各自资料。", "前端菜单按角色管理分组；学生详情含学籍与住宿历史入口；宿管详情显示负责楼栋；维修人员详情显示工种、楼栋范围、在岗和任务信息。"),
    ("阶段 4：楼栋数据要反映下级宿舍与床位", "楼栋总床位和空闲床位不显示或统计不正确；详情中下级宿舍只呈现为零散文本。", "楼栋床位数按楼下宿舍床位汇总，空闲床位根据未停用且未入住的床位统计；详情和编辑中以列表显示宿舍，允许新增和逐行编辑。", "ManagementService 对 room/bed/current assignment 做聚合；buildingRoomsPanel 使用表格、楼层筛选和每页 5 条；编辑可新增、编辑宿舍。"),
    ("阶段 5：宿舍与床位按真实空间层级查询", "楼层、宿舍筛选关系不清，列表列位、容量、空闲数和入住学生显示不准确。", "楼栋筛选联动楼层/宿舍；楼层再限制宿舍候选；床位筛选加入床位号、性别限制、学生姓名和床位状态；宿舍容量不再和床位数重复表达。", "rooms/beds 查询由 building_id、floor_no、room_id 联动；宿舍列表显示床位及空闲床位统计；床位列表展示楼栋、楼层、宿舍、床位号、状态、入住学生。"),
    ("阶段 6：入住调整从下拉选择改为可检索学生列表", "学生数量增加后，在床位下拉框或长页面中难以查找、分配学生。", "改成有姓名、学号、年级、专业等条件的学生列表，选中后填入床位；详情与编辑保持同一入口和流程。", "前端提供学生选择表格及分页；宿舍入住调整先进入独立调整区；床位调整有固定床位信息区和可分页学生表；保存变更通过入住分配服务落库。"),
    ("阶段 7：大列表改为单页分页与级联筛选", "后台查询记录多时页面过长，出现不希望的横向/下拉滚动，页码不便定位。", "管理员列表统一分页；支持输入页码跳转；下级列表一页 5 条；床位筛选按两行自适应布局。", "app.js 提供表格分页、页码输入、上一页/下一页；楼栋宿舍嵌套表格显式配置每页 5 条；测试脚本覆盖不同视口和分页交互。"),
    ("阶段 8：精细对齐和避免单字换行", "班级或当前住宿列会出现单字落到第二行；学号与姓名之间空白较大；操作按钮易换行且首按钮与表头不齐。", "按列重新分配宽度，把学号/姓名间多余空间用于当前住宿；超长值尽量均衡换行；操作按钮保持一行，首个按钮左对齐操作列表头。", "styles.css 针对学生列设置宽度与 text-wrap: balance；操作按钮 flex nowrap；角色操作容器 justify-content:flex-start；当前表格 CSS 和页面脚本已同步到 5173。"),
]
add_table(doc, ["迭代阶段", "起点/问题", "明确后的规则", "当前落地证据"], timeline, widths=[1.2, 1.9, 2.0, 2.0], font_size=8.0)

add_heading(doc, "3. 基本资料：楼栋、宿舍、床位", 1)
add_heading(doc, "3.1 楼栋", 2)
for text in [
    "查询列表不再仅展示楼栋静态字段，增加宿舍数、床位总数、空闲床位和已入住数等由下级资源计算出的汇总信息。统计基于 building → dorm_room → bed，并结合当前有效住宿分配判断占用状态。",
    "总床位来自该楼栋下实际建立的床位记录；空闲床位来自管理状态正常且没有当前入住分配的床位。容量是宿舍设计容量，不被误当成实际床位记录数。",
    "详情与编辑在楼栋基本资料同一页面展示楼栋宿舍列表。楼层筛选只影响该列表；按页显示 5 条，可翻页；编辑时可新增宿舍，也可在每行编辑宿舍基本资料，新增行的保存/取消与筛选工具对齐。",
    "宿舍字段包括所属楼栋、宿舍号、楼层、床位容量、性别限制。新增和修改均校验楼栋/宿舍关系及已建床位数量，避免容量小于现有床位。",
]:
    add_bullet(doc, text)

add_heading(doc, "3.2 宿舍", 2)
for text in [
    "宿舍查询列表按宿舍号、楼层、床位数、空闲床位数、性别限制等组织；床位数与容量含义重复时只保留一个列位，最终列表保留床位数，同时呈现空闲床位数。",
    "查询项支持楼栋、宿舍、楼层等层级筛选。选择楼栋后，楼层和宿舍候选被收窄；选择楼层后宿舍候选进一步收窄。性别限制也可作为筛选条件。",
    "详情和编辑都以床位/入住学生列表呈现宿舍内容，提供“调整入住学生”操作；实际选人进入独立调整界面，避免把整页做成长下拉。编辑状态下床位分配可以调整，取消和保存行为与详情中的调整入口保持一致。",
    "楼栋性别限制会约束宿舍性别；若有异性学生在住、容量不足或存在床位/维修等关联记录，服务层会拒绝不合业务规则的修改或删除。",
]:
    add_bullet(doc, text)

add_heading(doc, "3.3 床位", 2)
for text in [
    "床位列表定位到楼栋、楼层、宿舍、床位号，并展示管理状态与入住学生；学生姓名成为独立筛选条件，床位号用下拉选择。",
    "筛选分为两行：第一行楼栋/宿舍/楼层/床位；第二行性别限制/学生姓名/床位状态/查询与清空。宽度随视口变化，避免单项掉到第二行或形成横向滚动。",
    "床位状态区分为“可分配”“已入住”“停用”。已入住床位不再被描述为可分配；停用表示该床位不可继续使用。",
    "编辑床位时楼栋与宿舍使用分开的字段；可调整入住学生，入口与宿舍编辑一致。学生通过可搜索、可分页列表选择，确认保存后才写入住宿分配。",
]:
    add_bullet(doc, text)

add_heading(doc, "3.4 学生入住与调整的保存语义", 2)
add_p(doc, "住宿分配不是 student 表中的宿舍字段。当前住宿通过 accommodation_assignment 记录学生、床位、入住日期、退宿日期和分配类型；历史记录仍保留，当前住宿由视图按有效日期筛选。该设计让学生资料、床位资源与入住事件各自承担单一职责。")
add_table(doc, ["动作", "页面行为", "持久化结果"], [
    ("从学生新增/编辑打开宿舍调整", "离开表单进入分配界面；再次取消或关闭后返回，已填表单值保留为临时草稿", "尚未提交主表单前不写入学生资料或住宿分配"),
    ("在调整页保存选择", "回到原新增/编辑表单，当前住宿栏同步显示暂选楼栋、宿舍、床位", "仅更新草稿状态，不代表后端已保存"),
    ("提交学生新增/编辑主表单", "表单保存成功", "人员资料及本次住宿分配一并提交；实际业务保存点在此处"),
    ("从宿舍/床位详情调整入住", "固定资源摘要 + 可查询分页学生列表；保存后刷新对应详情", "调用住宿分配服务，记录有效分配及必要的退宿历史"),
], widths=[1.8, 2.7, 2.6], font_size=8.4)

add_heading(doc, "4. 角色管理：学生、宿管、维修人员", 1)
add_heading(doc, "4.1 菜单和统一操作", 2)
add_p(doc, "系统管理员菜单采用“角色管理”为一级入口，学生管理、宿管管理、维修人员管理为二级入口。每个页面按角色展示不同字段和过滤项，不把三类人员混用同一套资料栏。账号详情也嵌入对应人员详情，而不是让管理员只能看到登录账号。")
add_p(doc, "查询方式保留少量必要的多条件筛选，并提供一个可匹配账号、姓名、电话及学号/工号的关键词入口；界面通过提示说明可匹配字段，不把按钮或输入框标成“万能查询”。筛选栏控件统一高度，并根据窗口宽度换行。")
add_table(doc, ["页面", "列表信息", "详情/操作"], [
    ("学生管理", "学号、姓名、账号、学院、专业、年级、班级、当前住宿、账号状态", "查看个人资料、学籍、当前住宿和住宿历史；可进入报修/留宿记录；新增、编辑、启停、删除/批量删除、住宿分配"),
    ("宿管管理", "工号、姓名、账号、负责楼栋、在岗状态、账号状态", "查看基本资料和负责楼栋；配置楼栋范围；新增、编辑、启停、删除/批量删除"),
    ("维修人员管理", "工号、姓名、账号、工种、负责楼栋、在岗状态、账号状态", "查看人员资料、最大在办单量、工种和楼栋范围、维修历史/评价；配置范围；新增、编辑、启停、删除/批量删除"),
], widths=[1.2, 2.55, 3.35], font_size=8.2)

add_heading(doc, "4.2 学生档案", 2)
for text in [
    "学生基本字段落在 student：学生编号、10 位学号、姓名、性别、生日、联系电话、登录账号、密码摘要、启用状态、创建时间、更新时间。学籍状态由枚举表达，不用单独的“是否在读”布尔值重复表达休学、毕业、退学等情况。",
    "学院、专业、年级与班级不在每名学生行重复保存。student.class_id 指向 student_class；student_class 指向 major；major 指向 department，入学年份与年级由班级关联信息提供。这样班级变化不需要重复改写学生的学院和专业文本。",
    "筛选覆盖姓名、学号、楼栋/房间、学院、专业、年级、班级、账号状态和住宿状态等必要项；学生姓名和学号可分别查询。学院、专业、年级、班级候选按上下级选择联动。",
    "学生学号限制为 10 位数字；数据库有唯一约束和格式检查，服务端再次校验，并在编辑时排除当前学生后检查重复。",
]:
    add_bullet(doc, text)

add_heading(doc, "4.3 宿管档案", 2)
for text in [
    "宿管档案包含人员主键、登录账号、密码摘要、电话、启用状态、DM 前缀工号、姓名、性别、生日、入职时间、在岗状态、备注及创建/更新时间。工号统一为 DM 加 6 位数字。",
    "负责楼栋通过 manager_building_scope 关联，不把多个楼栋拼成宿管表中的重复文本字段；列表和详情通过查询组合展示。当前管理规则按楼栋负责人关系限制分配，系统需防止同一楼栋被重复配置给多个宿管。",
    "宿管列表提供姓名、工号、负责楼栋、账号状态和在岗状态等筛选；负责楼栋可从独立配置操作调整。",
]:
    add_bullet(doc, text)

add_heading(doc, "4.4 维修人员档案", 2)
for text in [
    "维修人员档案包含人员主键、账号凭据、电话、启用状态、RP 前缀工号、姓名、性别、生日、入职时间、在岗状态、最大在办工单数及时间字段。工号统一为 RP 加 6 位数字。",
    "负责楼栋和维修工种分别通过 repairer_building_scope、repairer_work_type 及其关系表保存，避免把多值关系冗余地写入维修人员主表。当前工单数、维修记录等属于业务数据，不作为档案冗余字段；详情页按需聚合展示。",
    "列表按姓名、工号、工种、负责楼栋、账号状态、在岗状态筛选；详情显示工种和楼栋范围，并提供维修历史与评价入口。",
]:
    add_bullet(doc, text)

add_heading(doc, "4.5 系统管理员个人资料", 2)
add_p(doc, "管理员本身不出现在普通角色管理的三个目录里，但个人资料按相同模型维护：SA 前缀工号、姓名、性别、生日、电话、入职时间、账号启用状态及时间字段。个人信息页面可编辑个人资料和修改密码；角色管理的重构不应删掉系统管理员自己的资料能力。")

add_heading(doc, "5. 数据模型与业务一致性", 1)
add_p(doc, "当前模型重点处理了三类容易重复或互相污染的数据：学生院系信息不再重复保存；住宿位置不写进学生档案；宿管/维修人员的多对多范围使用关联表。查询页显示需要的组合项时，通过视图或服务查询计算，不把显示结果重复写回基础表。")
add_table(doc, ["数据职责", "核心表/视图", "关系和约束"], [
    ("人员身份主键", "person", "提供 person_id；学生及三类员工档案以同一主键一对一扩展"),
    ("学生档案", "student", "学号 char(10) 唯一并检查 10 位数字；class_id 外键连接 student_class"),
    ("院系班级字典", "department → major → student_class", "学院不直接写在学生行；major 归属 department；班级归属专业并记录入学年份"),
    ("宿管档案/负责楼栋", "dorm_manager_profile + manager_building_scope", "宿管资料与 person 关联；楼栋范围独立关联，避免重复存楼栋文本"),
    ("维修人员档案/范围", "repairer_profile + repairer_building_scope + repairer_work_type", "员工资料、负责楼栋、维修工种各自分表/关系表表达"),
    ("管理员档案", "system_admin_profile", "管理员人员资料单独扩展 person；工号采用 SA 前缀规则"),
    ("住宿资源", "building → dorm_room → bed", "床位隶属于宿舍，宿舍隶属于楼栋；宿舍容量约束床位创建数量"),
    ("住宿历史", "accommodation_assignment + v_accommodation_current", "分配记录保存入住/退宿日期；当前住宿由视图筛选，不在 student 表冗余存放"),
    ("汇总显示", "v_student_profile、v_dorm_manager_profile、v_repairer_profile 等", "视图拼接档案和关联显示字段；楼栋/宿舍床位数由查询聚合计算"),
], widths=[1.25, 2.55, 3.35], font_size=8.1)

add_heading(doc, "5.1 已落实的格式与业务校验", 2)
for text in [
    "学生学号：数据库固定 char(10)、唯一键、正则格式约束；页面输入限制长度；服务端校验长度、数字和重复。",
    "员工工号：宿管 DM、维修人员 RP、系统管理员 SA 前缀，各自统一 6 位数字尾号，并有唯一与格式检查。",
    "宿舍容量：不能小于已建床位数；同楼栋宿舍号不能重复；宿舍性别限制受楼栋限制约束。",
    "床位状态：已有当前入住的床位不能作为可分配空床，也不能在不允许的条件下停用；删除受住宿或维修历史关联保护。",
    "入住分配：服务端检查床位状态、性别限制、目标床位占用与学生当前住宿；变更时结束旧分配并建立新的历史记录。",
]:
    add_bullet(doc, text)

add_heading(doc, "6. 页面布局与操作体验的反复修正", 1)
add_table(doc, ["反复出现的问题", "最终采用的处理规则", "当前实现留痕"], [
    ("某列过窄导致最后一个字单独换行", "按列分配宽度；不把一列的问题挪到相邻列；较长文本尽量均衡断行。", "角色表使用固定布局和分角色列宽；学生当前住宿列单独增加可用空间；班级列采用均衡换行规则。"),
    ("操作按钮落到第二行，首按钮与列名不齐", "操作按钮一行排列并禁用换行；操作区左对齐，使每行第一个按钮与“操作”表头左缘对齐。", "admin 表格操作区 flex nowrap / justify-content:flex-start；按钮不收缩。"),
    ("筛选项有的高、有的宽，布局时而拥挤时而空", "输入、下拉、操作按钮统一控制高度；按视口断点组织筛选网格；不缩小全局字体解决单一问题。", "统一 control height；床位筛选明确两行，断点在窄屏下调整每行字段数。"),
    ("列表太长或横向滚动", "长列表分页、可输入页码跳转；嵌套宿舍/学生列表使用单独分页；优先单页容纳，不增加多余下拉容器。", "通用 table pagination；楼栋宿舍每页 5 条；学生选择器分页。"),
    ("学号、姓名之间留白大，住宿列又换行", "重新分配列宽，把空白给当前住宿列，同时保持学号和姓名列可读。", "role-table-students 按列宽配置并针对常见宽度区间调整。"),
], widths=[1.65, 2.6, 2.9], font_size=8.2)
add_p(doc, "这些是功能性布局约束，不是全局缩字方案。普通界面字号保持可读，只在范围/数字筛选框这类局部位置按需要调整。")

add_heading(doc, "7. 当前验证证据与边界", 1)
add_heading(doc, "7.1 仓库中可核对的验证", 2)
add_table(doc, ["留痕项", "覆盖内容", "证据边界"], [
    ("role-management-smoke.cjs", "管理员登录、学生多条件/关键词查询、宿管和维修人员列表、管理员个人资料、移动端宽度", "仓库有测试脚本；browser-results.json 中的历史全系统结果日期为 2026-10-15，不能当作 10-30 本次执行结果"),
    ("admin-query-pagination-smoke.cjs", "楼栋床位统计、宿舍/床位筛选、5 条楼栋宿舍分页、入住调整、管理员分页及窄屏布局", "脚本包含明确断言；当前汇总未将其表述为本次重跑通过"),
    ("improvements-browser-smoke.cjs / document-improvements-smoke.cjs", "角色筛选、详情编辑、范围配置、学院/专业/班级级联、宿管楼栋约束和响应式表格", "脚本输出语句记录预期测试覆盖；需查看当次进程输出才能判断具体执行状态"),
    ("profile-v9-smoke.cjs", "角色菜单箭头、学生列表、学籍字典只读、四角色与移动端宽度", "有自动化脚本留痕；不能替代业务验收和数据库恢复验证"),
    ("当前 SQL 快照", "包含当前基础表、关联表、视图和示例数据；床位、楼栋、角色档案和学生等结构可直接核对", "快照反映导出时数据库；不能单独证明导入脚本在另一台机器执行结果一致"),
], widths=[1.85, 2.75, 2.55], font_size=8.0)

add_heading(doc, "7.2 不能由本报告推断的结论", 2)
for text in [
    "没有可靠 Git 提交历史可逐行定位每一次改动，因此阶段顺序按本对话提出需求的先后整理，不伪造提交号或每次修改日期。",
    "仓库包含历史测试结果和测试脚本；脚本存在不等于本次全部重新执行通过。报告只将可直接核对的脚本覆盖作为留痕。",
    "本报告范围只覆盖基本资料和角色管理，不代表整个四角色系统已完成所有需求验收。",
    "数据快照中使用的是开发演示数据。学院、专业、班级作为展示与筛选字典存在，不因此新增其独立维护菜单或增删改管理流程。",
]:
    add_bullet(doc, text)

add_heading(doc, "8. 关键代码与数据文件索引", 1)
add_table(doc, ["路径（相对项目根）", "用途"], [
    ("dorm-platform-frontend/src/management.js", "楼栋/宿舍/床位列表、详情、筛选联动、房间和床位入住调整、学生/宿管/维修人员管理与详情"),
    ("dorm-platform-frontend/src/app.js", "页面菜单、通用表格、分页与输入页码跳转"),
    ("dorm-platform-frontend/src/styles.css", "管理员列表列宽、操作列对齐、筛选控件和响应式布局"),
    ("dorm-platform-backend/src/main/java/com/dorm/platform/system/controller/ManagementController.java", "基本资料、住宿分配、角色目录、详情、资料编辑、批量删除等 HTTP 接口"),
    ("dorm-platform-backend/src/main/java/com/dorm/platform/system/service/ManagementService.java", "角色查询、楼栋/宿舍床位聚合、人员资料校验、住宿分配与资源业务规则"),
    ("sql/_localhost-2026_10_30_18_52_12-dump.sql", "当前导出的完整数据库结构和数据快照；含基础表、关联表和视图"),
    ("test-results/role-management-smoke.cjs", "角色管理浏览器交互测试脚本"),
    ("test-results/admin-query-pagination-smoke.cjs", "管理员查询、分页、住宿调整和布局测试脚本"),
    ("test-results/improvements-browser-smoke.cjs", "角色筛选、详情编辑、人员范围及布局测试脚本"),
    ("test-results/document-improvements-smoke.cjs", "学院专业班级级联、宿管楼栋约束及住宿分配流程测试脚本"),
], widths=[3.1, 4.05], font_size=8.1)

add_heading(doc, "结语", 1)
add_p(doc, "相较最初，基本资料已从单纯展示宿舍资源推进到可查看资源汇总、维护宿舍和床位、联动筛选并完成有记录的入住调整；角色管理已从围绕账号列表，推进到学生/宿管/维修人员各自的档案维护、详情、业务范围和数据校验。重复出现的列宽、换行、按钮错位和长列表问题也被整理成一致的页面规则，并落实到当前表格样式和分页交互中。")

doc.core_properties.title = "基本资料与角色管理改进总结及工作留痕"
doc.core_properties.subject = "宿舍事务平台系统管理员端"
doc.core_properties.author = "OpenAI Codex"
doc.core_properties.keywords = "宿舍管理,基本资料,角色管理,需求追踪,工作留痕"
zoom = doc.settings._element.find(qn("w:zoom"))
if zoom is None:
    zoom = OxmlElement("w:zoom")
    doc.settings._element.append(zoom)
zoom.set(qn("w:percent"), "100")
doc.save(OUT)
print(OUT)
